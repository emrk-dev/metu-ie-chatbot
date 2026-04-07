"""
ingest_forms.py
---------------
Extracts text from all files in the Forms/ folder and appends them
to the existing knowledge_base.json.

Supports: .pdf (pdfplumber), .docx (python-docx), .doc (macOS textutil)

Usage:
    cd scripts && python ingest_forms.py
"""

import json
import os
import subprocess
import tempfile
from pathlib import Path

import pdfplumber
from docx import Document

FORMS_DIR = Path(__file__).parent / "forms"
KB_FILE   = Path(__file__).parent / "knowledge_base.json"
MIN_CHUNK_LEN = 30


def extract_pdf(path: Path) -> list[str]:
    chunks = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if not text:
                continue
            for line in text.split("\n"):
                line = line.strip()
                if len(line) >= MIN_CHUNK_LEN:
                    chunks.append(line)
    return chunks


def extract_docx(path: Path) -> list[str]:
    doc = Document(path)
    chunks = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if len(text) >= MIN_CHUNK_LEN:
            chunks.append(text)
    # Also extract tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if len(row_text) >= MIN_CHUNK_LEN:
                chunks.append(row_text)
    return chunks


def extract_doc(path: Path) -> list[str]:
    """Use macOS textutil to convert .doc -> .txt, then read it."""
    with tempfile.TemporaryDirectory() as tmpdir:
        out_file = Path(tmpdir) / (path.stem + ".txt")
        result = subprocess.run(
            ["textutil", "-convert", "txt", "-output", str(out_file), str(path)],
            capture_output=True,
            text=True,
        )
        if result.returncode != 0:
            print(f"  [ERROR] textutil failed for {path.name}: {result.stderr}")
            return []
        if not out_file.exists():
            print(f"  [ERROR] No output file for {path.name}")
            return []
        raw = out_file.read_text(encoding="utf-8", errors="ignore")
        chunks = []
        for line in raw.split("\n"):
            line = line.strip()
            if len(line) >= MIN_CHUNK_LEN:
                chunks.append(line)
        return chunks


def build_entries(path: Path, chunks: list[str]) -> list[dict]:
    entries = []
    stem = path.stem.replace("-", "_").replace(" ", "_")
    for i, chunk in enumerate(chunks):
        entries.append({
            "id":    f"forms/{stem}_{i}",
            "url":   f"forms/{path.name}",
            "title": path.stem.replace("-", " ").replace("_", " ").title(),
            "text":  chunk,
        })
    return entries


def main():
    # Load existing knowledge base
    with open(KB_FILE, encoding="utf-8") as f:
        kb: list[dict] = json.load(f)

    existing_ids = {entry["id"] for entry in kb}
    print(f"Existing KB entries: {len(kb)}")

    new_entries = []

    for path in sorted(FORMS_DIR.iterdir()):
        suffix = path.suffix.lower()
        if suffix not in {".pdf", ".doc", ".docx"}:
            continue

        print(f"Processing: {path.name}")

        if suffix == ".pdf":
            chunks = extract_pdf(path)
        elif suffix == ".docx":
            chunks = extract_docx(path)
        elif suffix == ".doc":
            chunks = extract_doc(path)
        else:
            chunks = []

        entries = build_entries(path, chunks)
        # Skip duplicates
        entries = [e for e in entries if e["id"] not in existing_ids]
        new_entries.extend(entries)
        print(f"  -> {len(chunks)} chunks extracted, {len(entries)} new entries")

    kb.extend(new_entries)

    with open(KB_FILE, "w", encoding="utf-8") as f:
        json.dump(kb, f, ensure_ascii=False, indent=2)

    print(f"\nDone. Added {len(new_entries)} entries. Total KB size: {len(kb)}")


if __name__ == "__main__":
    main()
